"""Excel parsing: one page per worksheet, as a markdown table plus a structured table.

The workbook is opened twice (values and formulas) so formulas without a cached value
can be reported. Blocking: call from a worker thread. Never raises.
"""

from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

import openpyxl

from app.models.document import Page, PageKind, Table, table_to_markdown


@dataclass
class ExcelResult:
    pages: list[Page] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    error: str | None = None


def parse_excel(path: Path, row_cap: int) -> ExcelResult:
    result = ExcelResult()
    try:
        values_wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas_wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        result.error = f"cannot open workbook: {type(exc).__name__}: {exc}"
        return result

    try:
        sheets = zip(values_wb.worksheets, formulas_wb.worksheets, strict=True)
        for position, (values_ws, formulas_ws) in enumerate(sheets, start=1):
            try:
                page = _parse_sheet(values_ws, formulas_ws, position, row_cap)
            except Exception as exc:
                page = Page(
                    page_no=position,
                    kind=PageKind.SHEET,
                    sheet_name=values_ws.title,
                    warnings=[f"sheet could not be read: {exc}"],
                )
            if page is None:
                result.warnings.append(f"sheet '{values_ws.title}' is empty, skipped")
            else:
                result.pages.append(page)
    finally:
        values_wb.close()
        formulas_wb.close()
    return result


def _parse_sheet(values_ws, formulas_ws, position: int, row_cap: int) -> Page | None:
    # Read-only sheets can carry wrong stored dimensions; read everything instead.
    values_ws.reset_dimensions()
    formulas_ws.reset_dimensions()

    rows, truncated, missing_cached = _read_rows(values_ws, formulas_ws, row_cap)
    rows = _trim(rows)
    if not rows:
        return None

    warnings = []
    if getattr(values_ws, "sheet_state", "visible") != "visible":
        warnings.append("hidden sheet")  # TODO: confirm hidden-sheet policy with the team
    if truncated:
        warnings.append(f"truncated to the first {row_cap} rows")  # TODO: row cap placeholder
    if missing_cached:
        warnings.append(
            f"{missing_cached} formula cell(s) have no cached value; open and save in Excel"
        )

    table = Table(rows=rows, has_header=True)  # first non-empty row is taken as the header
    return Page(
        page_no=position,
        kind=PageKind.SHEET,
        sheet_name=values_ws.title,
        text=table_to_markdown(table),
        tables=[table],
        warnings=warnings,
    )


def _read_rows(values_ws, formulas_ws, row_cap: int) -> tuple[list[list[str]], bool, int]:
    """Rows as strings, starting at the first non-empty row. Returns (rows, truncated, missing)."""
    rows: list[list[str]] = []
    missing_cached = 0
    pairs = zip(
        values_ws.iter_rows(values_only=True), formulas_ws.iter_rows(values_only=True), strict=False
    )
    for values, formulas in pairs:
        missing_cached += sum(
            1
            for value, formula in zip(values, formulas, strict=False)
            if value is None and isinstance(formula, str) and formula.startswith("=")
        )
        cells = [_cell_text(value) for value in values]
        if not rows and not any(cells):
            continue  # leading empty rows
        if len(rows) == row_cap:
            if any(cells):
                return rows, True, missing_cached
            continue
        rows.append(cells)
    return rows, False, missing_cached


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop trailing empty rows and empty leading/trailing columns; pad rows to equal width.

    Merged cells need no handling: read-only mode returns the value only in the top-left cell.
    """
    while rows and not any(rows[-1]):
        rows.pop()
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    used = [col for col in range(width) if any(row[col] for row in rows)]
    return [row[used[0] : used[-1] + 1] for row in rows]


def _cell_text(value: object) -> str:
    """Numbers as stored, dates as ISO, strings with collapsed whitespace."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    return " ".join(str(value).split())
