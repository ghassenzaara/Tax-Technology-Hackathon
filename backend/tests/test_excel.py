import openpyxl

from app.ingest.excel import parse_excel
from app.models.document import PageKind
from tests.helpers import make_rich_excel


def test_one_page_per_sheet(tmp_path):
    result = parse_excel(make_rich_excel(tmp_path / "book.xlsx"), row_cap=100)

    assert result.error is None
    pages = {page.sheet_name: page for page in result.pages}
    assert list(pages) == ["Umsatz", "Merged", "Intern", "Formeln"]
    assert all(page.kind is PageKind.SHEET for page in result.pages)
    # page_no is the sheet's position, so the skipped empty sheet keeps number 2.
    assert [page.page_no for page in result.pages] == [1, 3, 4, 5]
    assert any("'Leer' is empty" in w for w in result.warnings)


def test_sheet_content_trim_and_dates(tmp_path):
    result = parse_excel(make_rich_excel(tmp_path / "book.xlsx"), row_cap=100)
    umsatz = result.pages[0]

    # Leading empty row and column are trimmed; dates are ISO; numbers stay as stored.
    assert umsatz.tables[0].rows == [
        ["Jahr", "Umsatz EUR", "Stichtag"],
        ["2023", "1234.5", "2023-12-31"],
    ]
    assert umsatz.text.splitlines()[0] == "| Jahr | Umsatz EUR | Stichtag |"
    assert umsatz.warnings == []


def test_merged_hidden_and_formula_warnings(tmp_path):
    pages = {p.sheet_name: p for p in parse_excel(make_rich_excel(tmp_path / "b.xlsx"), 100).pages}

    assert pages["Merged"].tables[0].rows[0] == ["Gesellschafter", ""]  # value only top-left
    assert pages["Intern"].warnings == ["hidden sheet"]
    assert "no cached value" in pages["Formeln"].warnings[0]


def test_row_cap_truncates_with_warning(tmp_path):
    workbook = openpyxl.Workbook()
    for i in range(10):
        workbook.active.append([f"row {i}"])
    workbook.save(tmp_path / "long.xlsx")

    [page] = parse_excel(tmp_path / "long.xlsx", row_cap=3).pages

    assert len(page.tables[0].rows) == 3
    assert "truncated" in page.warnings[0]


def test_corrupt_workbook_reports_error(tmp_path):
    (tmp_path / "bad.xlsx").write_bytes(b"PK\x03\x04 not really a zip")

    result = parse_excel(tmp_path / "bad.xlsx", row_cap=100)

    assert result.pages == [] and "cannot open workbook" in result.error
